from pathlib import Path

from docx import Document
from openpyxl import Workbook

HERE = Path(__file__).parent

doc = Document()
doc.add_heading("Положение об организационной структуре — ДО", 0)
doc.add_heading("Отдел закупок", level=1)
doc.add_paragraph("Отдел закупок обеспечивает проведение закупочных процедур, подготовку планов закупок и контроль исполнения договоров.")
doc.add_heading("Служба контроля", level=1)
doc.add_paragraph("Служба контроля осуществляет мониторинг исполнения договоров и готовит отчёты о соблюдении сроков.")
doc.save(HERE / "before.docx")

book = Workbook()
sheet = book.active
sheet.title = "Структура ПОСЛЕ"
sheet.append(["Подразделение", "Функция"])
sheet.append(["Отдел снабжения", "Обеспечивает проведение закупочных процедур и подготовку планов закупок."])
sheet.append(["Служба внутреннего аудита", "Осуществляет мониторинг исполнения договоров и готовит отчёты о соблюдении сроков."])
sheet.append(["Департамент цифровизации", "Организует разработку цифровых сервисов для автоматизации закупочного процесса."])
book.save(HERE / "after.xlsx")
