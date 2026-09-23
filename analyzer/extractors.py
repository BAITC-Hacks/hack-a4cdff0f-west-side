"""File extraction with stable document, page, sheet, and paragraph references."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path
import re

from docx import Document
import pymupdf
from openpyxl import load_workbook

from .models import ExtractedDocument, Segment, Source


def _segments_from_text(name: str, location: str, text: str) -> list[Segment]:
    result = []
    for chunk in re.split(r"\n+|(?<=[.!?])\s+(?=[А-ЯA-Z0-9])", text):
        clean = " ".join(chunk.split())
        if clean:
            result.append(Segment(Source(name, location, clean[:500]), clean))
    return result


def extract_document(name: str, content: bytes) -> ExtractedDocument:
    ext = Path(name).suffix.lower()
    document = ExtractedDocument(name=name, kind=ext.lstrip("."))
    if ext == ".docx":
        doc = Document(BytesIO(content))
        for idx, paragraph in enumerate(doc.paragraphs, 1):
            text = paragraph.text.strip()
            if text:
                location = f"абзац {idx}"
                document.segments.extend(_segments_from_text(name, location, text))
        for table_idx, table in enumerate(doc.tables, 1):
            for row_idx, row in enumerate(table.rows, 1):
                text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if text:
                    document.segments.extend(_segments_from_text(name, f"таблица {table_idx}, строка {row_idx}", text))
    elif ext == ".pdf":
        pdf = pymupdf.open(stream=content, filetype="pdf")
        for page_idx, page in enumerate(pdf, 1):
            document.segments.extend(_segments_from_text(name, f"страница {page_idx}", page.get_text()))
    elif ext in {".xlsx", ".xlsm"}:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        for sheet in workbook.worksheets:
            for row_idx, row in enumerate(sheet.iter_rows(values_only=True), 1):
                text = " | ".join(str(value).strip() for value in row if value is not None and str(value).strip())
                if text:
                    document.segments.extend(_segments_from_text(name, f"лист «{sheet.title}», строка {row_idx}", text))
        workbook.close()
    elif ext == ".doc":
        raise ValueError("Формат .doc не поддержан напрямую. Сохраните документ как .docx.")
    elif ext == ".xls":
        raise ValueError("Формат .xls не поддержан напрямую. Сохраните книгу как .xlsx.")
    else:
        raise ValueError(f"Неподдерживаемый формат: {ext or 'без расширения'}")
    if not document.segments:
        raise ValueError(f"В документе «{name}» не найден извлекаемый текст.")
    return document
